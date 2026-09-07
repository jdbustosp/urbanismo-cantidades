# comparar_idu.py - barrido de VALORES UNITARIOS del libro contra la
# base de precios IDU 2026-I Fase II (idu_apu_2026I.tsv + insumos).
# Uso:  python comparar_idu.py "<ruta del libro .xlsx>"
# Sale: comparacion_idu.tsv (una fila por precio del libro con el mejor
# match IDU, la relacion libro/IDU y bandera REVISAR cuando la
# desviacion supera el umbral). NO modifica el libro.
import sys, os, re, unicodedata
import openpyxl

AQUI = os.path.dirname(os.path.abspath(__file__))
LIBRO = sys.argv[1] if len(sys.argv) > 1 else None
if not LIBRO or not os.path.exists(LIBRO):
    print("Uso: python comparar_idu.py <libro.xlsx>"); sys.exit(1)

def norm(s):
    s = unicodedata.normalize("NFD", str(s or ""))
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    s = s.upper()
    s = re.sub(r"[^A-Z0-9 ]", " ", s)
    s = re.sub(r"\s+", " ", s).strip()
    return s

STOP = set("DE DEL LA EL LOS LAS Y E O U A EN CON PARA POR UN UNA AL INCLUYE TIPO SUMINISTRO INSTALACION SUM INST".split())

def words(s):
    return [w for w in norm(s).split() if w not in STOP and len(w) > 1]

# cargar referencia IDU (APU + insumos)
ref = []
for fn, kind in (("idu_apu_2026I.tsv", "APU"), ("idu_insumos_2026I.tsv", "INSUMO")):
    with open(os.path.join(AQUI, fn), encoding="utf-8") as f:
        hdr = f.readline().rstrip("\n").split("\t")
        for line in f:
            p = line.rstrip("\n").split("\t")
            if kind == "APU":
                nombre, um, val = p[4], p[5], p[6]
            else:
                nombre, um, val = p[3], p[4], p[5]
            try:
                v = float(val)
            except Exception:
                continue
            if v <= 0:
                continue
            ref.append((kind, nombre, um.strip().upper(), v, set(words(nombre))))
print("referencia IDU:", len(ref), "items")

# leer PRECIOS_UNITARIOS del libro: col B = nombre, col D = valor
wb = openpyxl.load_workbook(LIBRO, read_only=True, data_only=True)
if "PRECIOS_UNITARIOS" not in wb.sheetnames:
    print("el libro no tiene hoja PRECIOS_UNITARIOS"); sys.exit(1)
ws = wb["PRECIOS_UNITARIOS"]
libro = []
for row in ws.iter_rows(min_row=1, values_only=True):
    nombre = row[1] if len(row) > 1 else None
    valor = row[3] if len(row) > 3 else None
    if nombre and isinstance(valor, (int, float)) and valor >= 0:
        libro.append((str(nombre), float(valor)))
print("precios del libro:", len(libro))

def score(wa, wb_):
    if not wa or not wb_:
        return 0.0
    inter = len(wa & wb_)
    return inter / max(len(wa), 1) * 0.7 + inter / max(len(wb_), 1) * 0.3

out = os.path.join(AQUI, "comparacion_idu.tsv")
n_match = n_rev = 0
with open(out, "w", encoding="utf-8") as f:
    f.write("NOMBRE_LIBRO\tVU_LIBRO\tMATCH_IDU\tTIPO\tUM_IDU\tVU_IDU\tRATIO\tBANDERA\n")
    for nombre, vu in libro:
        wl = set(words(nombre))
        best = None; bs = 0.0
        for kind, rn, um, rv, wr in ref:
            s = score(wl, wr)
            # preferir APU (comparable con nuestros VU completos); un
            # INSUMO solo es material y castiga la comparacion
            if kind == "INSUMO":
                s *= 0.75
            if s > bs:
                bs = s; best = (kind, rn, um, rv)
        if best and bs >= 0.55:
            ratio = (vu / best[3]) if best[3] else 0.0
            aviso = " (vs INSUMO: IDU es solo material)" if best[0] == "INSUMO" else ""
            if vu == 0:
                flag = "SIN PRECIO"
            elif ratio < 0.45:
                flag = "REVISAR: MUY POR DEBAJO de IDU" + aviso
            elif ratio > 2.2:
                flag = "REVISAR: MUY POR ENCIMA de IDU" + aviso
            else:
                flag = "ok" + aviso
            if flag.startswith("REVISAR") or flag == "SIN PRECIO":
                n_rev += 1
            n_match += 1
            f.write("\t".join([nombre, str(vu), best[1], best[0], best[2],
                               str(best[3]), f"{ratio:.2f}", flag]) + "\n")
        else:
            f.write("\t".join([nombre, str(vu), "(sin par IDU claro)", "",
                               "", "", "", "sin referencia"]) + "\n")
print("matches:", n_match, "| para revisar:", n_rev)
print("salida:", out)
